"""Invoice, quotation and report documents (PDF + Excel)."""

import io
from collections.abc import Sequence
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from ..db import now_local
from ..models import CompanySettings, Invoice, Quotation
from ..serializers import INVOICE_STATUS_LABELS
from .pdf_base import (
    AMBER, CLAY, FROST, INK, INK_SOFT, MOSS, BrandPDF, logo_source, money, rupiah, t, trim_qty,
)

# The domestic invoice letterhead was fixed copy in the Laravel template.
COMPANY_LINES = [
    "PT Coffee Nation Prosperity",
    "Jl. Raya Bukit Kweni No 40, Sukodono, Kab. Sidoarjo, Jawa Timur",
    "info@vanillaroyal.id  |  +62 858-5366-9568  |  www.vanillaroyal.id",
]
METHOD_LABELS = {"transfer": "Transfer Bank", "qris": "QRIS"}


def _fmt_date(value: Any, with_time: bool = False) -> str:
    if not value:
        return "-"
    return value.strftime("%d %b %Y, %H:%M" if with_time else "%d %b %Y")


# ── Invoice ─────────────────────────────────────────────────────────────────


def invoice_pdf(invoice: Invoice) -> bytes:
    pdf = BrandPDF()
    pdf.footer_text = (
        "Vanilla Royal (PT Coffee Nation Prosperity) · info@vanillaroyal.id · +62 858-5366-9568"
        f"   ·   Dicetak: {_fmt_date(now_local(), True)}"
    )
    pdf.add_page()
    pdf.letterhead(
        logo_source(None, "CleanlLogo.png"), "Vanilla Royal", COMPANY_LINES,
        "Invoice", invoice.invoice_number,
        [INVOICE_STATUS_LABELS.get(invoice.status, invoice.status)],
    )

    top = pdf.get_y()
    pdf.eyebrow("Tagihan kepada")
    pdf.set_font("Times", "B", 12)
    pdf.cell(100, 6, t(invoice.customer_name), new_x="LMARGIN", new_y="NEXT")
    contact = [invoice.customer_phone, invoice.customer_email, invoice.customer_address]
    pdf.body("\n".join(c for c in contact if c), width=100)
    left_bottom = pdf.get_y()

    meta = [("Tanggal", _fmt_date(invoice.created_at))]
    if invoice.due_date:
        meta.append(("Jatuh Tempo", _fmt_date(invoice.due_date)))
    if invoice.paid_at:
        meta.append(("Lunas", _fmt_date(invoice.paid_at, True)))
    pdf.set_y(top)
    for label, value in meta:
        pdf.set_x(pdf.w - pdf.r_margin - 70)
        pdf.set_font("Courier", "B", 7)
        pdf.set_text_color(*CLAY)
        pdf.cell(30, 5, t(label.upper()))
        pdf.set_font("Helvetica", "", 8.5)
        pdf.set_text_color(*INK)
        pdf.cell(40, 5, t(value), align="R", new_x="LMARGIN", new_y="NEXT")
    pdf.set_y(max(left_bottom, pdf.get_y()) + 6)

    widths = [84, 18, 36, pdf.content_width - 138]
    pdf.items_table(
        ["Produk", "Qty", "Harga Satuan", "Total"], widths, ["L", "C", "R", "R"],
        [
            [(i.product_name, i.description or ""), (str(i.quantity), ""),
             (rupiah(i.unit_price), ""), (rupiah(i.total_price), "")]
            for i in invoice.items
        ],
    )
    pdf.ln(5)

    totals_x = pdf.w - pdf.r_margin - 78
    block_top = pdf.get_y()
    if invoice.notes:
        pdf.eyebrow("Catatan", CLAY)
        pdf.body(invoice.notes, width=pdf.content_width - 88)
    if invoice.status == "paid" and invoice.paid_at:
        # The stamp lives in the left column, level with the totals, so it never
        # overlaps figures or the confirmation block below.
        _paid_stamp(pdf, pdf.l_margin + 4, pdf.get_y() + (4 if invoice.notes else 2), invoice.paid_at)
    notes_bottom = pdf.get_y()
    pdf.set_y(block_top)
    rows = [("Subtotal", rupiah(invoice.subtotal))]
    rows += [(e.label, "+ " + rupiah(e.amount)) for e in invoice.extras]
    if invoice.discount and invoice.discount > 0:
        rows.append(("Diskon", "- " + rupiah(invoice.discount)))
    pdf.totals_box(rows, ("Total", rupiah(invoice.total_amount)), totals_x)
    if invoice.payment_method:
        pdf.set_x(totals_x)
        pdf.set_font("Helvetica", "", 8)
        pdf.set_text_color(*INK_SOFT)
        pdf.cell(78, 6, t("Metode: " + METHOD_LABELS.get(invoice.payment_method, "")), align="R",
                 new_x="LMARGIN", new_y="NEXT")
    pdf.set_y(max(notes_bottom, pdf.get_y()) + 6)

    if invoice.paid_at:
        _payment_record(pdf, invoice)
    return pdf.output_bytes()


def _paid_stamp(pdf: BrandPDF, x: float, y: float, paid_at: Any) -> None:
    """Double-ruled rubber-stamp mark, slightly tilted, in the settled (moss) colour."""
    w, h = 54.0, 22.0
    with pdf.rotation(-5, x=x + w / 2, y=y + h / 2):
        pdf.set_draw_color(*MOSS)
        pdf.set_line_width(0.8)
        pdf.rect(x, y, w, h, round_corners=True, corner_radius=1.5)
        pdf.set_line_width(0.25)
        pdf.rect(x + 1.4, y + 1.4, w - 2.8, h - 2.8, round_corners=True, corner_radius=1)
        pdf.set_text_color(*MOSS)
        pdf.set_font("Times", "B", 19)
        pdf.set_xy(x, y + 3.2)
        pdf.cell(w, 9, "LUNAS", align="C")
        pdf.set_font("Courier", "B", 6.8)
        pdf.set_xy(x, y + 12.6)
        pdf.cell(w, 4, t(f"PAID  ·  {paid_at:%d %b %Y}".upper()), align="C")
    pdf.set_line_width(0.2)
    pdf.set_text_color(*INK)
    pdf.set_y(y + h + 4)


def _payment_record(pdf: BrandPDF, invoice: Invoice) -> None:
    """Confirmation block: a moss rule with three labelled facts, no overlapping marks."""
    top = pdf.get_y()
    height = 15.0
    pdf.set_fill_color(230, 235, 224)
    pdf.rect(pdf.l_margin, top, pdf.content_width, height, style="F")
    pdf.set_fill_color(*MOSS)
    pdf.rect(pdf.l_margin, top, 1.2, height, style="F")

    facts = [
        ("Status pembayaran", "Lunas - terverifikasi"),
        ("Dikonfirmasi", _fmt_date(invoice.paid_at, True)),
        ("Metode", METHOD_LABELS.get(invoice.payment_method or "", "-")),
    ]
    col_w = (pdf.content_width - 8) / len(facts)
    for i, (label, value) in enumerate(facts):
        x = pdf.l_margin + 6 + i * col_w
        pdf.set_xy(x, top + 3)
        pdf.set_font("Courier", "B", 6.5)
        pdf.set_text_color(*MOSS)
        pdf.cell(col_w, 3.5, t(label.upper()))
        pdf.set_xy(x, top + 7.5)
        pdf.set_font("Helvetica", "B", 9)
        pdf.set_text_color(*INK)
        pdf.cell(col_w, 4.5, t(value))
    pdf.set_y(top + height + 4)


# ── Quotation ───────────────────────────────────────────────────────────────


def _item_specs(item: Any) -> str:
    parts = [item.species, item.grade, item.size, item.weight, item.moisture,
             item.packaging, item.condition, item.aroma]
    if item.country_of_origin:
        parts.append(f"Origin: {item.country_of_origin}")
    specs = " · ".join(p for p in parts if p)
    if item.hs_code:
        specs += ("\n" if specs else "") + f"HS Code: {item.hs_code}"
    if item.description:
        specs += ("\n" if specs else "") + item.description
    return specs


def _labelled_lines(pdf: BrandPDF, pairs: Sequence[tuple[str, Any]], width: float) -> None:
    for label, value in pairs:
        if not value:
            continue
        x = pdf.get_x()
        pdf.set_font("Helvetica", "B", 8)
        pdf.set_text_color(*INK)
        label_w = pdf.get_string_width(t(label) + ": ") + 1
        pdf.cell(label_w, 4.4, t(label) + ":")
        pdf.set_font("Helvetica", "", 8)
        pdf.set_text_color(*INK_SOFT)
        pdf.multi_cell(width - label_w, 4.4, t(value), new_x="LEFT", new_y="NEXT")
        pdf.set_x(x)
    pdf.set_text_color(*INK)


def quotation_pdf(quotation: Quotation, settings: CompanySettings) -> bytes:
    snap = quotation.pdf_company_snapshot or settings.to_snapshot()
    company = snap.get("company_name") or ""
    pdf = BrandPDF()
    pdf.show_page_numbers = True
    pdf.footer_text = (
        f"Issued by {company} · valid until {_fmt_date(quotation.valid_until)}"
        f" · generated {_fmt_date(now_local(), True)}"
    )
    pdf.add_page()
    address = (snap.get("address") or "") + (f", {snap['country']}" if snap.get("country") else "")
    contact = " | ".join(v for v in (snap.get("email"), snap.get("phone"), snap.get("website")) if v)
    pdf.letterhead(
        logo_source(snap.get("logo_data")), company, [address, contact],
        "International Export Quotation", quotation.quotation_number,
        [f"Date: {_fmt_date(quotation.quotation_date)}",
         f"Valid Until: {_fmt_date(quotation.valid_until)}"],
    )

    col_w = (pdf.content_width - 8) / 2
    top = pdf.get_y()
    pdf.eyebrow("Buyer")
    pdf.set_font("Times", "B", 12)
    pdf.multi_cell(col_w, 5.5, t(quotation.buyer_company_name), new_x="LMARGIN", new_y="NEXT")
    location = ", ".join(v for v in (quotation.buyer_city, quotation.buyer_state_province,
                                     quotation.buyer_postal_code, quotation.buyer_country) if v)
    lines = [quotation.buyer_contact_person, quotation.buyer_email, quotation.buyer_phone,
             quotation.buyer_address, location]
    if quotation.buyer_tax_vat_number:
        lines.append(f"Tax/VAT: {quotation.buyer_tax_vat_number}")
    pdf.body("\n".join(v for v in lines if v), size=8, width=col_w)
    left_bottom = pdf.get_y()

    right_x = pdf.l_margin + col_w + 8
    pdf.set_xy(right_x, top)
    pdf.set_font("Courier", "B", 7)
    pdf.set_text_color(*AMBER)
    pdf.cell(col_w, 4, "TRADE TERMS", new_x="LEFT", new_y="NEXT")
    incoterm = quotation.incoterm + (f" {quotation.incoterm_place}" if quotation.incoterm_place else "")
    _labelled_lines(pdf, [
        ("Incoterm", incoterm), ("Currency", quotation.currency),
        ("Payment Terms", quotation.payment_terms), ("MOQ", quotation.moq),
        ("Country of Origin", quotation.country_of_origin),
        ("Destination", quotation.destination_country),
    ], col_w)
    pdf.set_y(max(left_bottom, pdf.get_y()) + 5)

    shipping = [
        ("Port of Loading", quotation.port_of_loading),
        ("Port of Discharge", quotation.port_of_discharge),
        ("Final Destination", quotation.final_destination),
        ("Shipping Method", quotation.shipping_method),
        ("Production Lead Time", quotation.production_lead_time),
        ("Estimated Shipment", quotation.estimated_shipment),
    ]
    if any(v for _, v in shipping):
        pdf.eyebrow("Shipping information")
        _labelled_lines(pdf, shipping, pdf.content_width)
        pdf.ln(4)

    cur = quotation.currency
    widths = [92, 24, 30, pdf.content_width - 146]
    pdf.items_table(
        ["Product", "Qty", f"Unit Price ({cur})", f"Total ({cur})"], widths, ["L", "C", "R", "R"],
        [
            [(i.product_name, _item_specs(i)), (f"{trim_qty(i.quantity)} {i.unit}", ""),
             (money(i.unit_price), ""), (money(i.line_total), "")]
            for i in quotation.items
        ],
    )
    pdf.ln(5)

    totals_x = pdf.w - pdf.r_margin - 82
    block_top = pdf.get_y()
    if quotation.notes:
        pdf.eyebrow("Notes", CLAY)
        pdf.body(quotation.notes, size=8, width=pdf.content_width - 90)
    notes_bottom = pdf.get_y()
    pdf.set_y(block_top)
    rows = [("Subtotal", f"{cur} {money(quotation.subtotal)}")]
    optional = [
        ("Discount", quotation.discount_amount, "-"),
        ("Freight", quotation.freight_amount, "+"),
        ("Insurance", quotation.insurance_amount, "+"),
        (quotation.other_charges_label or "Other Charges", quotation.other_charges_amount, "+"),
    ]
    rows += [(label, f"{sign} {cur} {money(v)}") for label, v, sign in optional if v and v > 0]
    if quotation.tax_amount and quotation.tax_amount > 0:
        rate = f" ({trim_qty(quotation.tax_rate)}%)" if quotation.tax_rate else ""
        rows.append(((quotation.tax_label or "Tax") + rate, f"+ {cur} {money(quotation.tax_amount)}"))
    pdf.totals_box(rows, ("Grand Total", f"{cur} {money(quotation.grand_total)}"), totals_x, 82)
    pdf.set_y(max(notes_bottom, pdf.get_y()) + 6)

    if snap.get("terms_conditions"):
        pdf.add_page()
        pdf.eyebrow("Terms & Conditions")
        pdf.ln(1)
        pdf.body(snap["terms_conditions"], size=8)
        pdf.ln(6)

    bank_rows = [
        ("Bank Name", snap.get("bank_name")), ("Branch", snap.get("bank_branch")),
        ("SWIFT / BIC Code", snap.get("swift_bic")),
        ("Account Name (A/N)", snap.get("bank_account_name")),
        ("Account Number", snap.get("bank_account_number")),
        ("Currency Accepted", snap.get("bank_account_currency")),
        ("Bank Address", snap.get("bank_address")),
        ("Intermediary Bank", snap.get("intermediary_bank_name")),
        ("Intermediary SWIFT / BIC", snap.get("intermediary_bank_swift")),
        ("Intermediary Account No.", snap.get("intermediary_bank_account")),
        ("Charge Instruction", snap.get("payment_charge_instruction")),
    ]
    if snap.get("bank_name") or snap.get("bank_account_number"):
        pdf.eyebrow("Bank details for international wire transfer")
        pdf.ln(1)
        pdf.set_fill_color(*FROST)
        for label, value in [r for r in bank_rows if r[1]] + [("Payment Reference", quotation.quotation_number)]:
            pdf.set_font("Courier", "B", 7)
            pdf.set_text_color(*CLAY)
            pdf.cell(55, 6, "  " + t(label.upper()), fill=True)
            pdf.set_font("Helvetica", "B" if "SWIFT" in label else "", 8.5)
            pdf.set_text_color(*INK)
            pdf.multi_cell(pdf.content_width - 55, 6, t(value), new_x="LMARGIN", new_y="NEXT")
        pdf.ln(8)

    if pdf.get_y() > pdf.h - 60:
        pdf.add_page()
    sig_x = pdf.w - pdf.r_margin - 70
    pdf.set_x(sig_x)
    pdf.set_font("Helvetica", "", 8.5)
    pdf.cell(70, 5, "Authorized By,", new_x="LEFT", new_y="NEXT")
    pdf.ln(20)
    pdf.set_x(sig_x)
    pdf.set_draw_color(*INK)
    pdf.line(sig_x, pdf.get_y(), sig_x + 70, pdf.get_y())
    pdf.ln(1.5)
    pdf.set_x(sig_x)
    pdf.set_font("Times", "B", 11)
    pdf.cell(70, 5.5, t(snap.get("authorized_person_name") or company), new_x="LEFT", new_y="NEXT")
    pdf.set_font("Helvetica", "", 8)
    pdf.set_text_color(*INK_SOFT)
    if snap.get("authorized_person_position"):
        pdf.cell(70, 4.4, t(snap["authorized_person_position"]), new_x="LEFT", new_y="NEXT")
    pdf.cell(70, 4.4, t(company), new_x="LEFT", new_y="NEXT")
    return pdf.output_bytes()


# ── Reports ─────────────────────────────────────────────────────────────────


def report_pdf(invoices: Sequence[Invoice], total_revenue: float) -> bytes:
    pdf = BrandPDF(orientation="L")
    pdf.footer_text = f"Vanilla Royal · Laporan Invoice · Dicetak {_fmt_date(now_local(), True)}"
    pdf.show_page_numbers = True
    pdf.add_page()
    pdf.eyebrow("Vanilla Royal Payment")
    pdf.display("Laporan Invoice", 20)
    pdf.ln(2)
    pdf.body(f"{len(invoices)} invoice  ·  Total pendapatan (lunas): {rupiah(total_revenue)}", 9)
    pdf.ln(2)
    pdf.rule()
    pdf.ln(4)
    headers = ["No", "No. Invoice", "Customer", "Telepon", "Total", "Metode", "Status", "Tanggal", "Lunas"]
    widths = [10, 36, 54, 32, 34, 26, 34, 18, pdf.content_width - 244]
    aligns = ["C", "L", "L", "L", "R", "L", "L", "C", "C"]
    rows = []
    for n, inv in enumerate(invoices, start=1):
        rows.append([
            (str(n), ""), (inv.invoice_number, ""), (inv.customer_name, ""),
            (inv.customer_phone or "-", ""), (rupiah(inv.total_amount), ""),
            (METHOD_LABELS.get(inv.payment_method or "", "-"), ""),
            (INVOICE_STATUS_LABELS.get(inv.status, inv.status), ""),
            (inv.created_at.strftime("%d/%m/%Y") if inv.created_at else "-", ""),
            (inv.paid_at.strftime("%d/%m/%Y") if inv.paid_at else "-", ""),
        ])
    pdf.items_table(headers, widths, aligns, rows)
    return pdf.output_bytes()


def report_excel(invoices: Sequence[Invoice]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Laporan Invoice"
    headers = ["No", "No. Invoice", "Customer", "Telepon", "Total", "Metode", "Status", "Tanggal", "Lunas"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FBFAF8")
        cell.fill = PatternFill("solid", fgColor="1A1210")
    for n, inv in enumerate(invoices, start=1):
        ws.append([
            n, inv.invoice_number, inv.customer_name, inv.customer_phone or "-",
            float(inv.total_amount or 0), inv.payment_method or "-",
            INVOICE_STATUS_LABELS.get(inv.status, inv.status),
            inv.created_at.strftime("%d/%m/%Y") if inv.created_at else "-",
            inv.paid_at.strftime("%d/%m/%Y") if inv.paid_at else "-",
        ])
    for column, width in zip("ABCDEFGHI", (5, 20, 28, 16, 16, 12, 22, 12, 12)):
        ws.column_dimensions[column].width = width
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
